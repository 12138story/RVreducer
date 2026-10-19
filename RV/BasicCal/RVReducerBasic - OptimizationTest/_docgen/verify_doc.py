# -*- coding: utf-8 -*-
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import docx

path = r"RV减速器关键部件设计原理（基于C++计算引擎）.docx"
d = docx.Document(path)
paras = [p for p in d.paragraphs if p.text.strip()]
print("段落数(非空):", len(paras))
print("内嵌图片数:", len(d.inline_shapes))
print("表格数:", len(d.tables))
print("---- 开头非空段落 ----")
for p in paras[:6]:
    print(repr(p.text[:70]))
print("---- 图注 ----")
for p in paras:
    t = p.text.strip()
    if t.startswith("图 ") and len(t) < 60:
        print("CAP:", t)
print("---- 公式段示例(含(1-x)/(2-x)/(3-x)编号) ----")
cnt = 0
for p in paras:
    t = p.text.strip()
    if t.startswith("(") and ("-1)" in t or "-2)" in t or "-3)" in t):
        print("F:", t[:70])
        cnt += 1
        if cnt >= 14:
            break
print("---- 章节标题 ----")
for p in paras:
    t = p.text.strip()
    if t.startswith("第 ") or t.startswith("绪") or t.startswith("附录"):
        print("H1:", t)
    elif len(t) < 40 and (t[0].isdigit() and "." in t[:3] and "　" in t[:5]):
        print("  H2:", t)
